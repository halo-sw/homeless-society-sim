"""Seoul Metro congestion table loader (stage 8, Part B).

Reads ``data/raw/traffic/서울교통공사_지하철혼잡도정보_20260930.xlsx``
(Seoul Metro lines 1–8; sheets 평일/토요일/일요일; rows per line, station
number, station name and direction; 30-minute slot columns 05:30~06:00 …
24:30~25:00; values are average congestion %, 34 % ≈ all seats taken,
100 % = rated capacity) into ``data/build/region/congestion.json``:
per line / station / direction / slot, plus station numbers (used for
상선/하선 derivation in pipeline.network).

Licence: 공공누리 3유형 (출처표시+변경금지). Values are used for local
computation and display with attribution
"서울교통공사 지하철혼잡도정보 (2026-09-30 기준)". The xlsx and this
derived JSON live under data/ (gitignored) and are never committed or
redistributed; the browser fetches the JSON from the dev server at
runtime (never bundled).
"""

from __future__ import annotations

import json
import re
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_XLSX = (REPO_ROOT / "data" / "raw" / "traffic"
                / "서울교통공사_지하철혼잡도정보_20260930.xlsx")
DEFAULT_BUILD_DIR = REPO_ROOT / "data" / "build"
CONGESTION_FILE = "congestion.json"

ATTRIBUTION = "서울교통공사 지하철혼잡도정보 (2026-09-30 기준)"
LICENSE = "공공누리 3유형 (출처표시+변경금지) — 로컬 계산·표시용, 재배포 금지"

DAY_SHEETS = {"평일": "weekday", "토요일": "saturday", "일요일": "sunday"}
HEADER_FIXED = ["호선", "역번호", "역명", "상하구분"]
SLOT_RE = re.compile(r"^(\d{1,2}):(\d{2})~(\d{1,2}):(\d{2})$")


def slot_start_min(label: str) -> int:
    """'08:00~08:30' -> 480 (minutes since midnight; 24:xx past midnight)."""
    m = SLOT_RE.match((label or "").strip())
    if not m:
        raise ValueError(f"bad slot label {label!r}")
    return int(m.group(1)) * 60 + int(m.group(2))


def slot_for_minute(minute: int, slots: list[int]) -> int:
    """Index of the 30-minute slot containing `minute` (clamped)."""
    idx = 0
    for i, start in enumerate(slots):
        if minute >= start:
            idx = i
        else:
            break
    return idx


def load_workbook(xlsx: Path | None = None) -> dict:
    """Parse the xlsx into {days, station_numbers, slots, meta}."""
    import openpyxl

    xlsx = Path(xlsx) if xlsx else DEFAULT_XLSX
    if not xlsx.exists():
        raise FileNotFoundError(f"congestion xlsx not found: {xlsx}")
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    missing = [s for s in DAY_SHEETS if s not in wb.sheetnames]
    if missing:
        raise ValueError(f"xlsx misses sheets {missing} (has {wb.sheetnames})")

    days: dict[str, dict] = {}
    station_numbers: dict[str, dict[str, int]] = {}
    slots: list[int] | None = None
    slot_labels: list[str] = []
    for sheet, kind in DAY_SHEETS.items():
        ws = wb[sheet]
        rows = ws.iter_rows(values_only=True)
        header = [str(c).strip() if c is not None else "" for c in next(rows)]
        if header[:4] != HEADER_FIXED:
            raise ValueError(f"{sheet}: header wants {HEADER_FIXED}, got {header[:4]}")
        labels = header[4:]
        starts = [slot_start_min(c) for c in labels]
        if slots is None:
            slots, slot_labels = starts, labels
        elif starts != slots:
            raise ValueError(f"{sheet}: slot columns differ from 평일")
        table: dict[str, dict[str, dict[str, list]]] = {}
        for row in rows:
            line, num, name, direction = row[0], row[1], row[2], row[3]
            if line is None or name is None or direction is None:
                continue
            line = str(line).strip()
            name = str(name).strip()
            direction = str(direction).strip()
            values = [round(float(v), 1) if v is not None else None
                      for v in row[4:4 + len(starts)]]
            table.setdefault(line, {}).setdefault(name, {})[direction] = values
            try:
                station_numbers.setdefault(line, {})[name] = int(num)
            except (TypeError, ValueError):
                pass
        # Deterministic key order.
        days[kind] = {line: {name: {d: table[line][name][d]
                                        for d in sorted(table[line][name])}
                               for name in sorted(table[line])}
                      for line in sorted(table)}
    wb.close()
    return {
        "days": days,
        "station_numbers": {line: dict(sorted(stations.items()))
                            for line, stations in sorted(station_numbers.items())},
        "slots": slots or [],
        "slot_labels": slot_labels,
        "meta": {
            "source": ATTRIBUTION,
            "license": LICENSE,
            "lines": sorted({line for day in days.values() for line in day}),
        },
    }


def lookup(table: dict, day_kind: str, line: str, station: str,
           direction: str, minute: int) -> float | None:
    """Congestion % for a ride starting at `minute`, or None when unknown."""
    try:
        values = table["days"][day_kind][line][station][direction]
    except KeyError:
        return None
    idx = slot_for_minute(minute, table["slots"])
    if idx >= len(values):
        return None
    return values[idx]


def write_congestion(table: dict, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(table, indent=2, sort_keys=True,
                               ensure_ascii=False) + "\n", encoding="utf-8")
    return path


def run_build(xlsx: Path | None = None,
              build_dir: Path = DEFAULT_BUILD_DIR) -> dict:
    """Build ``data/build/region/congestion.json``; returns a summary."""
    table = load_workbook(xlsx)
    out = write_congestion(table, Path(build_dir) / "region" / CONGESTION_FILE)
    counts = {kind: sum(len(stations) for lines in day.values()
                        for stations in lines.values())
              for kind, day in table["days"].items()}
    return {"congestion": str(out), "lines": table["meta"]["lines"],
            "slots": len(table["slots"]), "rows": counts,
            "station_numbers": sum(len(s) for s in table["station_numbers"].values())}
